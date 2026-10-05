"""导出译文 Word 文档，原文图片按原顺序原样嵌入。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..core.model import Project, SegmentKind


def _set_font(run, name: str, size: int, bold: bool = False,
              color: tuple[int, int, int] | None = None) -> None:
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def export_word(project: Project, out_path: str | Path,
                style: str = "bilingual",
                asset_dir: str | Path | None = None,
                max_image_width: float = 6.0) -> Path:
    """导出 Word。

    style 取值：
        bilingual  原文与译文上下成对，适合校对
        target     只输出译文，适合交付
        source     只输出原文，适合备份
    """
    out_path = Path(out_path)
    doc = Document()

    base = doc.styles["Normal"]
    base.font.name = "Microsoft YaHei"
    base.font.size = Pt(11)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(title.add_run(project.name), "Microsoft YaHei", 18, bold=True)

    assets_root = Path(asset_dir) if asset_dir else (
        Path(project.asset_dir) if project.asset_dir else None
    )

    for seg in project.segments:
        if style in ("bilingual", "source") and seg.source.strip():
            _add_text(doc, seg.source, is_target=False,
                      size=12 if seg.kind == SegmentKind.CAPTION else 11)
        if style in ("bilingual", "target") and seg.target.strip():
            _add_text(doc, seg.target, is_target=True, size=12)
        if seg.kind == SegmentKind.CAPTION and style == "bilingual":
            _add_timecode(doc, seg)
        _insert_assets(doc, project, seg.index, assets_root, max_image_width)

    # 落在文末的图片
    _insert_assets(doc, project, len(project.segments), assets_root, max_image_width)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def _add_text(doc: Document, text: str, is_target: bool, size: int = 11) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    if is_target:
        p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run(text)
    color = (0x1F, 0x4E, 0x79) if is_target else (0x22, 0x22, 0x22)
    _set_font(run, "Microsoft YaHei", size, color=color)


def _add_timecode(doc: Document, seg) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"[{_fmt(seg.start)} - {_fmt(seg.end)}]")
    _set_font(run, "Consolas", 8, color=(0x88, 0x88, 0x88))


def _fmt(sec: float) -> str:
    m, s = divmod(int(sec), 60)
    return f"{m:02d}:{s:02d}"


def _insert_assets(doc: Document, project: Project, after_index: int,
                   assets_root: Path | None, max_width: float) -> None:
    ids = project.asset_anchors.get(str(after_index))
    if not ids or assets_root is None:
        return
    for asset_id in ids:
        asset = project.asset_by_id(asset_id)
        if asset is None:
            continue
        path = assets_root / asset.filename
        if not path.exists():
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width = min(max_width, asset.width / 96.0) if asset.width else max_width
        p.add_run().add_picture(str(path), width=Inches(width))
        if asset.caption:
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_font(cap.add_run(asset.caption), "Microsoft YaHei", 9,
                      color=(0x66, 0x66, 0x66))
